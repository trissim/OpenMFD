"""Focused checks for September dose mapping and non-destructive display cropping."""

import importlib
import json
import os
import sys
from pathlib import Path

import numpy as np
import pytest
from PIL import Image


@pytest.fixture(scope='module')
def figure():
    folder = Path(__file__).resolve().parents[2] / 'papers/mfd_platform/supplementary_bio'
    sys.path.insert(0, str(folder))
    return importlib.import_module('generate_september_axotomy_figure')


def test_description_owns_conditions(figure):
    mapping = figure.parse_description(
        ' 2-6 7-11_x000D_\nB-C: Y27 epob_x000D_\nD-E FC-A FC-A+EpoB_x000D_\nF-G DMSO DMSO'
    )
    assert len(mapping) == 60
    assert mapping['B', 2] == 'Y27'
    assert mapping['C', 11] == 'EpoB'
    assert mapping['D', 6] == 'FC-A'
    assert mapping['E', 7] == 'EpoB+FC-A'
    assert mapping['G', 11] == 'DMSO'
    assert [figure.DOSES[(col - 2) % 5] for col in range(2, 12)] == [40, 20, 10, 5, 0] * 2


def test_unrecognized_map_fails_closed(figure):
    with pytest.raises(ValueError):
        figure.parse_description('2-7 8-11\nB-C F04 FC-A')


def test_display_crops_keep_entire_image_rectangles(figure, tmp_path):
    pixels = np.full((1100, 1000, 3), 255, dtype=np.uint8)
    pixels[150:1090, 5:460] = [20, 35, 45]
    pixels[150:1090, 540:995] = [50, 15, 20]
    pixels[25:40, 250:750] = 0
    path = tmp_path / 'pair.png'
    Image.fromarray(pixels).save(path)
    crops, boxes = figure.extract_image_rectangles(path)
    assert boxes == [[5, 150, 460, 1090], [540, 150, 995, 1090]]
    np.testing.assert_array_equal(np.asarray(crops[0]), pixels[150:1090, 5:460])
    np.testing.assert_array_equal(np.asarray(crops[1]), pixels[150:1090, 540:995])


def test_real_workbook_doses_and_replicates(figure):
    path = figure.DEFAULT_ROOT / 'neurite outgrowth/results.xlsx'
    if not path.exists():
        pytest.skip('External September dataset is not mounted')
    wells, descriptions = figure.read_outgrowth(path)
    assert len(wells) == 120
    assert len(descriptions) == 2
    counts = wells.groupby(['plate', 'condition', 'nominal_dose_uM']).size()
    assert set(counts.values) == {2, 4}
    assert (counts.loc['24099', 'DMSO'] == 4).all()
    assert wells.set_index(['plate', 'well']).loc[('24098', 'B02'), 'nominal_dose_uM'] == 40
    assert wells.set_index(['plate', 'well']).loc[('24098', 'B06'), 'nominal_dose_uM'] == 0


def test_layout_has_no_overlap_or_schematic(figure):
    boxes = list(figure.BOXES.values())
    for i, (x, y, w, h) in enumerate(boxes):
        assert x >= 0 and y >= 0
        assert x + w <= figure.CANVAS[0] and y + h <= figure.CANVAS[1]
        for xx, yy, ww, hh in boxes[i + 1:]:
            assert x + w <= xx or xx + ww <= x or y + h <= yy or yy + hh <= y


def test_merged_dose_axis_includes_all_series_and_larger_second_plate(figure, monkeypatch, tmp_path):
    rows = []
    for plate, conditions, baseline in (
        ('24098', ('F04', '09-027', '06-049', '09-037', '08-115', '09-079'), 1),
        ('24099', ('DMSO', 'Y27', 'FC-A', 'EpoB', 'EpoB+FC-A'), 100),
    ):
        for condition in conditions:
            for dose in figure.DOSES:
                for replicate in (0, 1):
                    rows.append(dict(plate=plate, condition=condition,
                                     nominal_dose_uM=dose, value=baseline + dose / 10 + replicate))
    from matplotlib.figure import Figure
    captured = []
    def capture(fig, *args, **kwargs):
        assert len(fig.axes) == 1
        axis = fig.axes[0]
        doses = [d for d in figure.DOSES if d != 0]
        assert [label.get_text() for label in axis.get_xticklabels()] == [str(d) for d in doses] * 11
        conditions = [text.get_text() for text in axis.texts]
        assert len(conditions) == 11 and len(set(conditions)) == 11
        bars = fig.axes[0].patches
        assert len(bars) == 44
        from matplotlib.colors import to_rgba
        for i, bar in enumerate(bars):
            baseline = 1 if i // 4 < 6 else 100
            assert bar.get_height() == baseline + doses[i % 4] / 10 + 0.5
            np.testing.assert_allclose(bar.get_facecolor(), to_rgba(figure.COLORS[conditions[i // 4]]))
        assert len({bar.get_facecolor() for bar in bars}) == 11
        baselines = [line for line in axis.lines if line.get_linestyle() == ':']
        assert len(baselines) == 1
        np.testing.assert_allclose(baselines[0].get_ydata(), 102)
        np.testing.assert_allclose(baselines[0].get_xdata(), [0, 1])
        assert baselines[0].get_transform() == axis.get_yaxis_transform()
        captured.append(fig.axes[0].get_ylim())
    monkeypatch.setattr(Figure, 'savefig', capture)
    figure.dose_panel(figure.pd.DataFrame(rows), ('value', 'Test', 'Value', ''),
                      tmp_path / 'test.png', (5100, 810))
    assert captured
    assert all(low == 0 and high > 101 for low, high in captured)


def test_prepared_image_checksum_fails_closed(figure, tmp_path):
    prepared = tmp_path / 'prepared'
    prepared.mkdir()
    asset = prepared / 'image.png'
    Image.new('RGB', (40, 50), 'black').save(asset)
    record = dict(plate='P1', well='R02C02', source='original.tiff',
                  display_image='image.png', display_sha256=figure.digest(asset))
    (prepared / 'figure_manifest.json').write_text(json.dumps({'image_inventory': [record]}))
    metrics = figure.pd.DataFrame([dict(plate='P1', well='R02C02',
        metric='post_skeleton_pixels_per_pre_segmented_object', eligible=True, condition='DMSO')])
    index = figure.pd.DataFrame([dict(plate='P1', well='R02C02')])
    output = tmp_path / 'output'
    output.mkdir()
    _, inventory = figure.asset_inventory(metrics, index, output, prepared)
    assert figure.digest(output / inventory[0]['display_image']) == record['display_sha256']
    Image.new('RGB', (40, 50), 'red').save(asset)
    with pytest.raises(ValueError, match='checksum mismatch'):
        figure.asset_inventory(metrics, index, output, prepared)


def test_built_exports_match_source_measurements(figure):
    output = Path(os.environ.get('SEPTEMBER_FIGURE_OUTPUT',
        figure.BIO / 'figure_output_september_20260909'))
    if not (output / 'figure_manifest.json').exists():
        pytest.skip('Build outputs are not available')
    manifest = json.loads((output / 'figure_manifest.json').read_text())
    for filename, expected in manifest['source_files'].items():
        assert figure.digest(output / 'sources' / filename) == expected
    source = figure.pd.read_csv(output / 'sources/relative_half_axon_method_defined_metric_qc.csv')
    plotted = figure.pd.read_csv(output / 'axotomy_plotted_rows.csv')
    assert plotted.groupby('panel').size().to_dict() == {'B': 43, 'C': 31, 'D': 35, 'E': 31}
    expected = source[source.eligible & source.metric.isin([s[0] for s in figure.AXON_SPECS])]
    columns = ['plate', 'well', 'metric', 'analytical_value', 'condition']
    sort = ['plate', 'well', 'metric']
    figure.pd.testing.assert_frame_equal(
        plotted[columns].sort_values(sort).reset_index(drop=True),
        expected[columns].sort_values(sort).reset_index(drop=True))
    values = source.pivot(index=['plate', 'well'], columns='metric', values='analytical_value')
    ratio = values['post_skeleton_pixels_per_pre_segmented_object'].dropna()
    np.testing.assert_allclose(ratio, (values.skeleton_pixels / values.pre_objects).loc[ratio.index])
    wells, _ = figure.read_outgrowth(output / 'sources/results.xlsx')
    exported = figure.pd.read_csv(output / 'outgrowth_wells.csv', dtype={'plate': str})
    figure.pd.testing.assert_frame_equal(wells, exported)
    from statistics import mean, stdev
    summary = figure.pd.read_csv(output / 'axotomy_summary.csv').set_index(['panel', 'condition'])
    for key, group in plotted.groupby(['panel', 'condition']):
        row = summary.loc[key]
        assert row['count'] == len(group)
        np.testing.assert_allclose([row['mean'], row['std']],
                                   [mean(group.analytical_value), stdev(group.analytical_value)])
    growth_summary = figure.pd.read_csv(output / 'outgrowth_summary.csv', dtype={'plate': str})
    for r in growth_summary.itertuples():
        group = wells[(wells.plate == r.plate) & (wells.condition == r.condition)
                      & (wells.nominal_dose_uM == r.nominal_dose_uM)][r.metric]
        assert r.count == len(group)
        np.testing.assert_allclose([r.mean, r.std], [mean(group), stdev(group)])
    excel_rows = figure.pd.read_excel(output / 'September_PI_data.xlsx', sheet_name='Axotomy plotted')
    figure.pd.testing.assert_frame_equal(plotted, excel_rows, check_dtype=False)
    selected = figure.pd.read_csv(output / 'representative_selection.csv')
    assert set(selected.condition) == set(figure.CONDITIONS)
    for r in selected.itertuples():
        assert values.loc[(r.plate, r.well), [s[0] for s in figure.AXON_SPECS]].notna().all()
    for record in manifest['image_inventory']:
        assert figure.digest(output / record['display_image']) == record['display_sha256']
    assert len(manifest['image_inventory']) == 48
    assert sum(r['source'] is not None for r in manifest['image_inventory']) == 35


def test_powerpoint_contains_all_panels_without_cropping(figure):
    from pptx import Presentation
    output = Path(os.environ.get('SEPTEMBER_FIGURE_OUTPUT',
        figure.BIO / 'figure_output_september_20260909'))
    path = output / 'September_axotomy_and_outgrowth_editable.pptx'
    if not path.exists():
        pytest.skip('PowerPoint is not built')
    presentation = Presentation(path)
    assert len(presentation.slides) == 1
    shapes = presentation.slides[0].shapes
    assert len(shapes) == 20
    labels = []
    for shape in shapes:
        assert shape.left >= 0 and shape.top >= 0
        assert shape.left + shape.width <= presentation.slide_width
        assert shape.top + shape.height <= presentation.slide_height
        if shape.has_text_frame:
            labels.append(shape.text)
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    assert run.font.name == 'Arial' and run.font.size.pt == 11
        else:
            assert all(getattr(shape, 'crop_' + edge) == 0 for edge in ('left', 'right', 'top', 'bottom'))
            w, h = shape.image.size
            assert abs((shape.width / shape.height) / (w / h) - 1) < 1e-5
    assert labels == list('ABCDEFGHIJ')
