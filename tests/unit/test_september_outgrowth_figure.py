"""Provenance and scope checks for the outgrowth-only September figure."""

import importlib
import json
import sys
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from PIL import Image
from pptx import Presentation


@pytest.fixture(scope="module")
def figure():
    folder = Path(__file__).resolve().parents[2] / "papers/mfd_platform/supplementary_bio"
    sys.path.insert(0, str(folder))
    return importlib.import_module("generate_september_outgrowth_figure")


def test_layout_has_six_images_and_four_outgrowth_plots_without_overlap(figure):
    assert tuple(figure.BOXES) == tuple("ABCDEFGHIJ")
    assert figure.BOXES["A"] == (150, 90, 1600, 1750)
    assert figure.BOXES["D"] == (150, 1940, 1600, 1750)
    for earlier, later in zip("GHI", "HIJ", strict=True):
        first, second = figure.BOXES[earlier], figure.BOXES[later]
        assert first[3] >= 950
        assert 0 <= second[1] - first[1] - first[3] <= 30
    boxes = list(figure.BOXES.values())
    for index, (x, y, width, height) in enumerate(boxes):
        assert x >= 0 and y >= 0
        assert x + width <= figure.CANVAS[0]
        assert y + height <= figure.CANVAS[1]
        for xx, yy, ww, hh in boxes[index + 1:]:
            assert x + width <= xx or xx + ww <= x or y + height <= yy or yy + hh <= y


def test_display_window_is_shared_across_complete_images(figure):
    first = np.full((256, 256), 200, dtype=np.uint16)
    second = np.full((256, 256), 220, dtype=np.uint16)
    first[20:70, 20:70] = 1000
    second[50:100, 50:100] = 1500
    originals = [first.copy(), second.copy()]
    expected = np.percentile(np.concatenate([first.ravel(), second.ravel()]), 99.9)
    assert figure.pooled_uint16_percentile([first, second], 99.9) == \
        pytest.approx(expected)
    rendered, display = figure.display_stack([first, second])
    assert len(rendered) == 2
    assert all(image.size == (256, 256) for image in rendered)
    assert display["raw_window"][0] == 200
    assert display["rolling_ball_radius_native_px"] == 192
    assert display["background_subtraction_fraction"] == pytest.approx(0.70)
    assert display["post_low_percentile"] == 35
    assert display["post_low_value"] < display["post_high_value"]
    assert display["highlight_ceiling_fraction"] == pytest.approx(1.0)
    assert all(np.asarray(image)[..., 1].max() <= 255 for image in rendered)
    dapi, dapi_display = figure.dapi_stack([first, second])
    assert dapi_display["window"][0] == pytest.approx(220)
    assert len(dapi) == 2
    assert figure.overlay_dapi(rendered[0], dapi[0]).size == (256, 256)
    np.testing.assert_array_equal(first, originals[0])
    np.testing.assert_array_equal(second, originals[1])


def test_built_figure_has_condition_verified_images_and_filtered_metaxpress_graphs(figure):
    output = figure.DEFAULT_OUTPUT
    if not (output / "figure_manifest.json").is_file():
        pytest.skip("Outgrowth figure has not been built")
    manifest = json.loads((output / "figure_manifest.json").read_text())
    assert manifest["quantitative_source"].startswith("MetaXpress results.xlsx")
    assert len(manifest["images"]) == 6
    assert [(row["condition"], row["dose_um"]) for row in manifest["images"]] == [
        ("DMSO", None), ("F04", 20), ("06-049", 40), ("08-115", 5),
        ("Y27", 40), ("FC-A", 40),
    ]
    assert manifest["plotted_conditions"]["24099"] == ["DMSO", "Y27", "FC-A"]
    assert manifest["excluded_conditions"] == ["EpoB", "EpoB+FC-A"]
    plotted = pd.read_csv(output / "outgrowth_plotted_wells.csv")
    assert len(plotted) == 100
    assert not {"EpoB", "EpoB+FC-A"} & set(plotted.condition)
    assert "axon_method" not in manifest
    assert figure.digest(figure.DEFAULT_WORKBOOK) == manifest["workbook"]["sha256"]
    assert figure.digest(output / manifest["workbook"]["retained_copy"]) == \
        manifest["workbook"]["sha256"]
    for record in manifest["images"]:
        assert record["identity_check"] == \
            "original and coded site-1 FITC and DAPI pixels equal"
        assert record["original_site_shape_yx"] == [1024, 1024]
        assert record["source_pixel_size_um_yx"] == [1.3556, 1.3556]
        assert record["scale_bar"]["label_um"] == 500
        assert record["scale_bar"]["effective_length_um"] == pytest.approx(500, abs=2)
        assert record["display"]["scope"].startswith("six complete stitched FITC mosaics")
        assert record["display"]["raw_window"] == manifest["images"][0]["display"]["raw_window"]
        assert record["display"]["post_high_value"] == \
            manifest["images"][0]["display"]["post_high_value"]
        assert record["display"]["post_low_value"] == \
            manifest["images"][0]["display"]["post_low_value"]
        assert record["dapi_display"]["window"] == \
            manifest["images"][0]["dapi_display"]["window"]
        assert record["condition_outline_color"] == \
            figure.september.COLORS[record["condition"]]
        assert record["condition_outline_width_px"] == 24
        assert 0 < record["presentation_qc"]["green_percentiles_25_50_75_95"][1] < 30
        assert 0 < record["presentation_qc"]["green_ceiling_fraction"] < 0.01
        assert 0 < record["presentation_qc"]["blue_above_60_fraction"] < 0.05
        for source in (record["original_site_1_fitc"], record["coded_site_1_fitc"],
                       record["original_site_1_dapi"], record["coded_site_1_dapi"],
                       record["stitched_fitc"], record["stitched_dapi"]):
            assert figure.digest(Path(source["path"])) == source["sha256"]
    for label in "GHIJ":
        with Image.open(output / "panels" / f"panel_{label}.png") as panel:
            assert panel.size == figure.BOXES[label][2:]
    assert (output / "outgrowth_wells.csv").read_bytes() == \
        (figure.BIO / "figure_output_september_20260909/outgrowth_wells.csv").read_bytes()
    presentation = Presentation(output / "September_outgrowth_stitched_editable.pptx")
    assert len(presentation.slides) == 1
    assert len(presentation.slides[0].shapes) >= 20
    assert any(shape.name == "Figure note" and shape.text == figure.FIGURE_NOTE
               for shape in presentation.slides[0].shapes)
