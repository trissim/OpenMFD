#!/usr/bin/env python3
"""September axotomy and independent neurite-outgrowth dose-series figure.

Run with the project .venv. Source results are never modified. The workbook's
description cells own treatment placement; doses were supplied by the user.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import zipfile
from importlib.metadata import version
from pathlib import Path
from types import SimpleNamespace

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageOps

import generate_axotomy_supplementary_figure as shared
import generate_combined_axotomy_supplementary_figure as assembly
from axotomy_plate_mosaic import save_plate_mosaic


BIO = Path(__file__).resolve().parent
DEFAULT_ROOT = Path('/run/media/ts/0BA20E780BA20E78/axotomy_testing/2026-09-04-(complete)')
CONDITIONS = ('DMSO', 'FC-A', 'F04', 'EpoB', 'EpoB+FC-A')
DOSES = (40, 20, 10, 5, 0)
COLORS = {
    'DMSO': '#64717D', 'FC-A': '#00856A', 'F04': '#C88710',
    'EpoB': '#C44C63', 'EpoB+FC-A': '#805CA5', 'Y27': '#2679B2',
    '09-027': '#239FA7', '06-049': '#A24C72', '09-037': '#527B37',
    '08-115': '#454F9C', '09-079': '#C4542E',
}
AXON_SPECS = (
    ('pre_objects', 'Pre-axotomy CTB-positive cell count', 'Segmented objects'),
    ('skeleton_pixels', 'Regenerating axon measurement', 'Skeleton pixels'),
    ('dapi_objects', 'Endpoint nuclear count', 'Segmented objects'),
    ('post_skeleton_pixels_per_pre_segmented_object',
     'Regenerating axon measurement / PRE CTB cells', 'Skeleton pixels / PRE segmented object'),
)
OUTGROWTH_SPECS = (
    ('mean_outgrowth', 'Mean neurite outgrowth per cell', 'Length (export units)',
     'Mean Outgrowth Per Cell (Neurite Outgrowth)'),
    ('median_process_length', 'Median process length', 'Length (export units)',
     'Cell: Median Process Length (Neurite Outgrowth)'),
    ('branches_per_cell', 'Branching per cell', 'Branches / cell',
     'Mean Branches Per Cell (Neurite Outgrowth)'),
    ('cell_count', 'Cell count', 'Cells (exported summary)',
     'Number of Cells (Neurite Outgrowth)'),
)
BOXES = {
    'A': (150, 90, 2500, 1720),
    'B': (2800, 90, 2450, 830),
    'C': (2800, 990, 2450, 830),
    'D': (150, 1930, 2500, 930),
    'E': (2800, 1930, 2450, 930),
    'F': (150, 2970, 5100, 1100),
    **{letter: (150, 4200 + i * 880, 5100, 810)
       for i, letter in enumerate('GHIJ')},
}
CANVAS = (5400, 7780)


def digest(path: Path) -> str:
    with path.open('rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest()


def canonical_condition(name: str) -> str:
    return {'epob': 'EpoB', 'FC-A+EpoB': 'EpoB+FC-A'}.get(name, name)


def parse_description(description: str) -> dict[tuple[str, int], str]:
    lines = description.replace('_x000D_', '').strip().splitlines()
    blocks = re.findall(r'(\d+)-(\d+)', lines[0])
    if blocks != [('2', '6'), ('7', '11')]:
        raise ValueError(f'Unexpected column declaration: {lines[0]}')
    mapping = {}
    for line in lines[1:]:
        match = re.fullmatch(r'\s*([A-H])-([A-H]):?\s+(\S+)\s+(\S+)\s*', line)
        if not match:
            raise ValueError(f'Unparsed description row: {line!r}')
        first, last, left, right = match.groups()
        for row in range(ord(first), ord(last) + 1):
            for (start, stop), condition in zip(blocks, (left, right), strict=True):
                for col in range(int(start), int(stop) + 1):
                    key = (chr(row), col)
                    if key in mapping:
                        raise ValueError(f'Duplicate plate-map position {key}')
                    mapping[key] = canonical_condition(condition)
    if len(mapping) != 60:
        raise ValueError(f'Expected 60 assigned wells, found {len(mapping)}')
    return mapping


def read_outgrowth(path: Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    records, descriptions = [], []
    workbook = load_workbook(path, read_only=True, data_only=True)
    plate_name = plate_id = description = headers = mapping = None
    try:
        for sheet in workbook:
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                key = row[0]
                if key == 'Plate Name':
                    plate_name, headers, mapping = row[1], None, None
                elif key == 'Plate ID':
                    plate_id = str(row[1])
                elif key == 'Description':
                    description = row[1]
                    mapping = parse_description(description)
                    descriptions.append(dict(plate=plate_id, plate_name=plate_name,
                                             description=description))
                elif key is None and 'Number of Cells (Neurite Outgrowth)' in row:
                    headers = list(row)
                elif isinstance(key, str) and re.fullmatch(r'[A-H]\d{2}', key):
                    if all(v is None for v in row[1:]):
                        continue
                    if mapping is None or headers is None:
                        raise ValueError(f'Missing metadata at Excel row {row_number}')
                    col = int(key[1:])
                    condition = mapping[(key[0], col)]
                    position = (col - 2) % 5
                    record = dict(plate=plate_id, plate_name=plate_name, well=key,
                                  condition=condition, column=col, row=key[0],
                                  block='left' if col <= 6 else 'right',
                                  series_position=position + 1,
                                  nominal_dose_uM=DOSES[position],
                                  excel_sheet=sheet.title, excel_row=row_number)
                    for metric, _, _, field in OUTGROWTH_SPECS:
                        value = row[headers.index(field)]
                        if not isinstance(value, (int, float)) or not np.isfinite(value):
                            raise ValueError(f'Invalid {field} at {plate_name} {key}')
                        record[metric] = value
                    records.append(record)
    finally:
        workbook.close()
    frame = pd.DataFrame(records)
    if len(frame) != 120 or frame.duplicated(['plate', 'well']).any():
        raise ValueError('Expected exactly 120 unique outgrowth wells')
    return frame, pd.DataFrame(descriptions)


def longest_run(mask: np.ndarray) -> tuple[int, int]:
    boundaries = np.diff(np.r_[False, mask, False].astype(int))
    runs = list(zip(np.flatnonzero(boundaries == 1), np.flatnonzero(boundaries == -1)))
    if not runs:
        raise ValueError('No image rectangle detected')
    return max(runs, key=lambda pair: pair[1] - pair[0])


def extract_image_rectangles(path: Path) -> tuple[list[Image.Image], list[list[int]]]:
    """Remove only the white Matplotlib layout and captions, keeping both image rectangles."""
    with Image.open(path) as source:
        image = source.convert('RGB')
    pixels = np.asarray(image)
    height, width = pixels.shape[:2]
    nonwhite = np.min(pixels, axis=2) < 245
    mid = nonwhite[height // 3: 2 * height // 3].mean(axis=0) > 0.60
    bounds = np.diff(np.r_[False, mid, False].astype(int))
    columns = [(a, b) for a, b in zip(np.flatnonzero(bounds == 1),
                                    np.flatnonzero(bounds == -1)) if b - a > width * 0.2]
    if len(columns) != 2:
        raise ValueError(f'Expected two image rectangles in {path}; found {columns}')
    boxes, crops = [], []
    for x0, x1 in columns:
        y0, y1 = longest_run(nonwhite[:, x0:x1].mean(axis=1) > 0.60)
        if y1 - y0 < height * 0.75:
            raise ValueError(f'Image body unexpectedly short in {path}')
        box = [int(x0), int(y0), int(x1), int(y1)]
        boxes.append(box)
        crops.append(image.crop(box))
    return crops, boxes


def make_pair_asset(source: Path, target: Path) -> list[list[int]]:
    crops, boxes = extract_image_rectangles(source)
    height = 1600
    resized = [ImageOps.contain(crop, (1600, height), Image.Resampling.LANCZOS)
               for crop in crops]
    width = sum(crop.width for crop in resized) + 12
    canvas = Image.new('RGB', (width, height), 'white')
    x = 0
    for crop in resized:
        canvas.paste(crop, (x, (height - crop.height) // 2))
        x += crop.width + 12
    canvas.save(target)
    return boxes


def asset_inventory(metrics: pd.DataFrame, index: pd.DataFrame, output: Path,
                    prepared_inputs: Path | None = None):
    assets = output / 'assets'
    assets.mkdir(exist_ok=True)
    rows, inventory = [], []
    paired = {(r.plate, r.well): r for r in index.itertuples()}
    prepared = {}
    if prepared_inputs is not None:
        manifest = json.loads((prepared_inputs / 'figure_manifest.json').read_text())
        prepared = {(r['plate'], r['well']): r for r in manifest['image_inventory']}
    for (plate, well), data in metrics.groupby(['plate', 'well'], sort=True):
        source = paired.get((plate, well))
        target = assets / f'{plate}_{well}.png'
        flags = data.set_index('metric')
        excluded = not bool(flags.loc['post_skeleton_pixels_per_pre_segmented_object', 'eligible'])
        if source is not None:
            if prepared_inputs is not None:
                record = prepared[(plate, well)].copy()
                asset = prepared_inputs / record['display_image']
                if digest(asset) != record['display_sha256']:
                    raise ValueError(f'Prepared image checksum mismatch: {asset}')
                if asset.resolve() != target.resolve():
                    shutil.copy2(asset, target)
            else:
                source_path = Path(source.overlay)
                if not source_path.exists():
                    # Image indices preserve the original drive path on a new machine.
                    parts = source_path.parts
                    offset = parts.index('openhcs_reanalysis_20260908')
                    source_path = index.attrs['source_root'].joinpath(*parts[offset:])
                boxes = make_pair_asset(source_path, target)
                record = dict(plate=plate, well=well, source=str(source_path),
                              source_sha256=digest(source_path), image_rectangles=boxes)
            record.update(display_image=str(target.relative_to(output)),
                          display_sha256=digest(target))
            inventory.append(record)
        else:
            placeholder = Image.new('RGB', (1200, 1600), '#EFF1F3')
            draw = ImageDraw.Draw(placeholder)
            reason = 'Incomplete POST' if bool(data.post_available.iloc[0]) else 'No POST acquisition'
            draw.multiline_text((600, 730), reason.replace(' ', '\n', 1),
                                font=shared.pil_font(65), fill='#626B74',
                                anchor='mm', align='center', spacing=12)
            placeholder.save(target)
            inventory.append(dict(plate=plate, well=well, source=None,
                                  display_image=str(target.relative_to(output)),
                                  display_sha256=digest(target), reason=reason))
        rows.append(SimpleNamespace(plate=plate.lower(), well=well,
                                    row=int(well[1:3]), column=int(well[4:6]),
                                    condition=data.condition.iloc[0], path=target,
                                    excluded=excluded and source is not None))
    return rows, inventory


def representative_selection(metrics: pd.DataFrame, index: pd.DataFrame) -> pd.DataFrame:
    values = metrics.pivot(index=['plate', 'well', 'condition'], columns='metric',
                           values='analytical_value').reset_index()
    fields = [spec[0] for spec in AXON_SPECS]
    candidates = values.dropna(subset=fields).merge(index[['plate', 'well']], on=['plate', 'well'])
    selected = []
    for condition in CONDITIONS:
        group = candidates[candidates.condition == condition].copy()
        if group.empty:
            raise ValueError(f'No QC-eligible representative for {condition}')
        matrix = group[fields]
        center = matrix.median()
        mad = (matrix - center).abs().median().replace(0, 1)
        group['selection_distance'] = (((matrix - center) / mad) ** 2).sum(axis=1)
        selected.append(group.sort_values(['selection_distance', 'plate', 'well']).iloc[0])
    return pd.DataFrame(selected).reset_index(drop=True)


def representatives_panel(selected: pd.DataFrame, output: Path, path: Path):
    width, height = BOXES['F'][2:]
    canvas = Image.new('RGB', (width, height), 'white')
    draw = ImageDraw.Draw(canvas)
    draw.text((width / 2, 4), 'Representative axotomy devices',
              font=shared.pil_font(46, bold=True), fill=shared.INK, anchor='ma')
    tile_width = width // len(selected)
    for i, row in enumerate(selected.itertuples()):
        x = i * tile_width
        draw.text((x + tile_width / 2, 75), row.condition,
                  font=shared.pil_font(46, bold=True), fill=COLORS[row.condition], anchor='ma')
        draw.text((x + tile_width / 2, 135), 'PRE CTB             POST axons',
                  font=shared.pil_font(46), fill=shared.INK, anchor='ma')
        source = output / 'assets' / f'{row.plate}_{row.well}.png'
        image = shared.image_thumbnail(source, (tile_width - 20, height - 200))
        canvas.paste(image, (x + 10, 195))
    canvas.save(path, dpi=(300, 300))


def dose_panel(frame: pd.DataFrame, spec: tuple, path: Path, size: tuple[int, int],
               order: dict[str, tuple[str, ...]] | None = None,
               dose_note: str | None = None,
               font_scale: float = 1.0,
               show_note: bool = True):
    metric, title, unit, _ = spec
    fig, axis = plt.subplots(figsize=(size[0] / 300, size[1] / 300), dpi=300)
    if order is None:
        order = {
            '24098': ('F04', '09-027', '06-049', '09-037', '08-115', '09-079'),
            '24099': ('DMSO', 'Y27', 'FC-A', 'EpoB', 'EpoB+FC-A'),
        }
    bar_doses = tuple(dose for dose in DOSES if dose != 0)
    offsets = np.linspace(-0.30, 0.30, len(bar_doses))
    conditions_in_order = [condition for conditions in order.values() for condition in conditions]
    group = 0
    for plate, conditions in order.items():
        for condition in conditions:
            data = frame[(frame.plate == plate) & (frame.condition == condition)]
            stats = data.groupby('nominal_dose_uM')[metric].agg(['mean', 'std'])
            for dose, offset in zip(bar_doses, offsets, strict=True):
                x = group + offset
                values = data.loc[data.nominal_dose_uM == dose, metric]
                axis.bar(x, stats.loc[dose, 'mean'], width=0.175, color=COLORS[condition],
                         edgecolor=shared.INK, linewidth=0.35, zorder=2)
                axis.errorbar(x, stats.loc[dose, 'mean'], yerr=stats.loc[dose, 'std'],
                              fmt='none', ecolor=shared.INK, capsize=1.5,
                              elinewidth=0.7, capthick=0.7, zorder=3)
                jitter = np.linspace(-0.032, 0.032, len(values))
                axis.scatter(x + jitter, values, color=shared.INK,
                             s=8, alpha=0.75, zorder=4)
            axis.annotate(condition, xy=(group, 0), xycoords=axis.get_xaxis_transform(),
                          xytext=(0, -25), textcoords='offset points',
                          ha='center', va='top', fontsize=11 * font_scale,
                          color=COLORS[condition])
            group += 1
    control = frame.loc[frame.condition == 'DMSO', metric]
    if control.empty or not np.isfinite(control).all():
        raise ValueError(f'Missing or invalid DMSO reference for {metric}')
    axis.axhline(control.mean(), color=shared.INK, linestyle=':', linewidth=1.25, zorder=5)
    axis.set_xticks([group + offset for group in range(len(conditions_in_order)) for offset in offsets],
                    [str(dose) for _ in conditions_in_order for dose in bar_doses])
    axis.tick_params(axis='x', length=2, pad=3, labelsize=10 * font_scale)
    axis.tick_params(axis='y', labelsize=10 * font_scale)
    axis.set_xlim(-0.6, len(conditions_in_order) - 0.4)
    axis.spines[['top', 'right']].set_visible(False)
    axis.grid(axis='y', color=shared.GRID, linewidth=0.6)
    axis.set_axisbelow(True)
    axis.set_ylim(bottom=0)
    axis.set_ylabel(unit, fontsize=11 * font_scale)
    fig.suptitle(f'Neurite outgrowth: {title}', y=0.99,
                 fontsize=12 * font_scale)
    fig.subplots_adjust(left=0.07, right=0.99,
                        bottom=0.32 if show_note else 0.23, top=0.90)
    if show_note:
        fig.text(0.5, 0.025,
                 dose_note or 'Dose (uM); bars: mean +/- SD; points: wells; dotted line: pooled DMSO mean.',
                 ha='center', fontsize=11 * font_scale, color=shared.MUTED)
    fig.savefig(path, dpi=300)
    fig.savefig(path.with_suffix('.svg'))
    fig.savefig(path.with_suffix('.pdf'))
    plt.close(fig)


def build(root: Path, output: Path, prepared_inputs: Path | None = None) -> None:
    corrected = root / 'openhcs_reanalysis_20260908/results/treatment_corrected_20260909'
    metric_path = corrected / 'method_defined_metrics/relative_half_axon_method_defined_metric_qc.csv'
    index_path = corrected / 'paired_figure_index.csv'
    workbook_path = root / 'neurite outgrowth/results.xlsx'
    if prepared_inputs is not None:
        prepared_manifest = json.loads((prepared_inputs / 'figure_manifest.json').read_text())
        metric_path = prepared_inputs / 'sources' / metric_path.name
        index_path = prepared_inputs / 'sources' / index_path.name
        workbook_path = prepared_inputs / 'sources' / workbook_path.name
        for path in (metric_path, index_path, workbook_path):
            if digest(path) != prepared_manifest['source_files'][path.name]:
                raise ValueError(f'Prepared source checksum mismatch: {path}')
    output.mkdir(parents=True, exist_ok=True)
    panels = output / 'panels'
    panels.mkdir(exist_ok=True)
    sources = output / 'sources'
    sources.mkdir(exist_ok=True)
    report = (prepared_inputs / 'sources' if prepared_inputs else corrected) / 'REPORT.md'
    for path in (metric_path, index_path, workbook_path, report):
        if path.resolve() != (sources / path.name).resolve():
            shutil.copy2(path, sources / path.name)
    metrics = pd.read_csv(metric_path)
    if metrics.duplicated(['plate', 'well', 'metric']).any():
        raise ValueError('Duplicate axotomy measurements')
    if len(metrics) != 336 or set(metrics.condition) != set(CONDITIONS):
        raise ValueError('Unexpected September metric inventory')
    if metrics.loc[~metrics.eligible, 'analytical_value'].notna().any():
        raise ValueError('Ineligible measurements must remain missing')
    index = pd.read_csv(index_path)
    index.attrs['source_root'] = root
    outgrowth, descriptions = read_outgrowth(workbook_path)
    outgrowth.to_csv(output / 'outgrowth_wells.csv', index=False)
    descriptions.to_csv(output / 'workbook_plate_descriptions.csv', index=False)
    print('Validated corrected axotomy data and 120 dose-mapped outgrowth wells', flush=True)
    records, inventory = asset_inventory(metrics, index, output, prepared_inputs)
    print('Prepared all 35 image pairs and missing-acquisition placeholders', flush=True)
    paths = {label: panels / f'panel_{label}.png' for label in BOXES}
    save_plate_mosaic(records, paths['A'], plates=('p1', 'p2'),
                      image_path_for_record=lambda record: record.path,
                      is_excluded=lambda record: record.excluded,
                      plate_colors=shared.PLATE_COLORS, legend_plates=('p1', 'p2'),
                      title='September axotomy: both plates',
                      font_regular=shared.FONT_REGULAR, font_bold=shared.FONT_BOLD,
                      legend_panel_range='B-E', excluded_label='excluded ratio',
                      subtitle='PRE CTB | POST axons; grey: no complete pair',
                      label_font_px=44)
    plotted = []
    for label, (metric, title, unit) in zip('BCDE', AXON_SPECS, strict=True):
        data = metrics[(metrics.metric == metric) & metrics.eligible].copy()
        plotted.append(data.assign(panel=label))
        plot_records = [SimpleNamespace(plate=row.plate.lower(), condition=row.condition,
                                       row=int(row.well[1:3]), column=int(row.well[4:6]),
                                       outlier=False, value=float(row.analytical_value))
                        for row in data.itertuples()]
        shared.save_metric_panel(plot_records,
                                 shared.MetricSpec(panel=label, value='value', title=title, ylabel=unit), paths[label],
                                 BOXES[label][2:], conditions=CONDITIONS, plates=('p1', 'p2'),
                                 condition_colors={key: COLORS[key] + '55' for key in CONDITIONS},
                                 condition_labels={'EpoB+FC-A': 'EpoB + FC-A'})
    selected = representative_selection(metrics, index)
    selected.to_csv(output / 'representative_selection.csv', index=False)
    representatives_panel(selected, output, paths['F'])
    for label, spec in zip('GHIJ', OUTGROWTH_SPECS, strict=True):
        dose_panel(outgrowth, spec, paths[label], BOXES[label][2:])
    exact = pd.concat(plotted, ignore_index=True)
    axon_summary = exact.groupby(['panel', 'condition']).analytical_value.agg(['count', 'mean', 'std']).reset_index()
    long = outgrowth.melt(id_vars=[c for c in outgrowth.columns if c not in {s[0] for s in OUTGROWTH_SPECS}],
                          value_vars=[s[0] for s in OUTGROWTH_SPECS], var_name='metric', value_name='value')
    growth_summary = long.groupby(['plate', 'condition', 'nominal_dose_uM', 'metric']).value.agg(['count', 'mean', 'std']).reset_index()
    exact.to_csv(output / 'axotomy_plotted_rows.csv', index=False)
    axon_summary.to_csv(output / 'axotomy_summary.csv', index=False)
    growth_summary.to_csv(output / 'outgrowth_summary.csv', index=False)
    with pd.ExcelWriter(output / 'September_PI_data.xlsx', engine='openpyxl') as writer:
        for name, frame in (
            ('Axotomy all QC', metrics), ('Axotomy plotted', exact),
            ('Axotomy summary', axon_summary), ('Outgrowth wells', outgrowth),
            ('Outgrowth summary', growth_summary), ('Plate descriptions', descriptions),
            ('Representatives', selected),
        ):
            frame.to_excel(writer, sheet_name=name, index=False)
            sheet = writer.sheets[name]
            sheet.freeze_panes = 'A2'
            sheet.auto_filter.ref = sheet.dimensions
            for col in sheet.columns:
                sheet.column_dimensions[col[0].column_letter].width = min(55, max(16, len(str(col[0].value)) + 2))
    stem = output / 'September_axotomy_and_outgrowth'
    assembly.compose_figure(paths, stem.with_suffix('.png'), panel_boxes=BOXES, canvas_size=CANVAS)
    assembly.build_powerpoint(paths, output / f'{stem.name}_editable.pptx',
                              panel_boxes=BOXES, canvas_size=CANVAS)
    manifest = {
        'sources': {str(p): digest(p) for p in (metric_path, index_path, workbook_path)},
        'source_files': {p.name: digest(p) for p in (metric_path, index_path, workbook_path)},
        'original_sources': (prepared_manifest.get('original_sources', prepared_manifest['sources'])
                             if prepared_inputs else {str(p): digest(p) for p in (metric_path, index_path, workbook_path)}),
        'doses_left_to_right_uM': DOSES, 'dose_authority': 'User declaration, 2026-09-09',
        'positive_control': 'FC-A, user declaration; no extra label in figure',
        'panels': {k: list(v) for k, v in BOXES.items()}, 'canvas_pixels': CANVAS,
        'axon_method': 'relative_half_axon; original PRE and nuclear measurement methods',
        'representative_rule': 'Nearest condition median across the four axotomy metrics, MAD-scaled squared distance, eligible pairs only; tie by plate/well',
        'image_processing': 'Existing approved paired TIFF pixels; remove only white layout/captions and resize without stretching. No new contrast or segmentation.',
        'image_inventory': inventory,
        'outgrowth_fields': {spec[0]: spec[3] for spec in OUTGROWTH_SPECS},
        'outgrowth_length_units': 'Workbook does not declare measurement units; retained as export units. Raw TIFF spatial calibration is um, but export-unit confirmation is pending.',
        'outgrowth_aggregation': 'Use each workbook well value as exported. No pooling across doses, no new exclusions, no individual-cell reconstruction.',
        'warnings': ['Two independent assay datasets; axotomy drug doses not supplied.',
                     'DMSO nominal dose is a matched dilution position, not DMSO concentration.',
                     'Combination dose is nominal as supplied; component concentrations not specified.',
                     'Cell: Median Process Length is an exported well summary, not a newly computed median over pooled cells.',
                     'Cell counts are fractional exported summaries, not total cells in the well.',
                     'No donor replication or inferential significance is asserted.'],
    }
    (output / 'figure_manifest.json').write_text(json.dumps(manifest, indent=2) + '\n')
    print(f'Wrote {stem}.png/.pdf, editable PPTX, and September_PI_data.xlsx', flush=True)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source-root', type=Path, default=DEFAULT_ROOT)
    parser.add_argument('--output-dir', type=Path, default=BIO / 'figure_output_september_20260909')
    parser.add_argument('--prepared-inputs', type=Path,
                        help='Previous output directory with checksum-verified sources and prepared images; no external drive needed')
    parser.add_argument('--package', action='store_true',
                        help='Also create a ZIP with figures, data, prepared images and rebuild code')
    args = parser.parse_args()
    build(args.source_root.resolve(), args.output_dir.resolve(),
          args.prepared_inputs.resolve() if args.prepared_inputs else None)
    if args.package:
        write_handoff(args.output_dir.resolve())


def write_handoff(output: Path) -> Path:
    """Bundle display inputs and the shared rendering code, not the raw acquisition archive."""
    path = output.with_name('September_axotomy_and_outgrowth_PI_handoff.zip')
    root = BIO.parents[2]
    code = [Path(__file__), Path(shared.__file__), Path(assembly.__file__),
            Path(assembly.control.__file__), BIO / 'axotomy_plate_mosaic.py',
            root / 'tests/unit/test_september_axotomy_figure.py']
    files = {f'figure/{p.relative_to(output)}': p for p in sorted(output.rglob('*')) if p.is_file()}
    files.update({f'code/{p.relative_to(root)}': p for p in code})
    files['README.md'] = BIO / 'SEPTEMBER_AXOTOMY_AND_OUTGROWTH.md'
    packages = ('matplotlib', 'numpy', 'pandas', 'openpyxl', 'Pillow', 'python-pptx', 'pytest')
    requirements = '\n'.join(f'{name}=={version(name)}' for name in packages) + '\n'
    checksums = {name: digest(source) for name, source in files.items()}
    checksums['requirements.txt'] = hashlib.sha256(requirements.encode()).hexdigest()
    with zipfile.ZipFile(path, 'w', compression=zipfile.ZIP_DEFLATED, compresslevel=6) as archive:
        for name, source in files.items():
            archive.write(source, name)
        archive.writestr('requirements.txt', requirements)
        archive.writestr('SHA256SUMS.json', json.dumps(checksums, indent=2) + '\n')
    print(f'Wrote {path}', flush=True)
    return path


if __name__ == '__main__':
    main()
