#!/usr/bin/env python3
"""Build a compact drug-response plus control-reproducibility figure."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

import generate_axotomy_control_supplementary_figure as control
import generate_axotomy_supplementary_figure as drug
from PIL import Image, ImageDraw, ImageFont, ImageOps
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

BIO_DIR = Path(__file__).resolve().parent
OUTPUT_DIR = BIO_DIR / "figure_output_combined"
FINAL_DPI = 300
CANVAS_WIDTH = 5400
CANVAS_HEIGHT = 7900
INK = "#17212B"

PANEL_BOXES = {
    "A": (140, 80, 1390, 1370),
    "B": (1650, 80, 3610, 1370),
    "C": (140, 1530, 2510, 1250),
    "D": (2780, 1530, 2480, 1250),
    "E": (140, 2860, 2510, 1250),
    "F": (2780, 2860, 2480, 1250),
    "G": (140, 4190, 5120, 930),
    "H": (140, 5200, 2510, 1250),
    "I": (2780, 5200, 2480, 1250),
    "J": (140, 6530, 2510, 1250),
    "K": (2780, 6530, 2480, 1250),
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--control-asset-root",
        type=Path,
        required=True,
        help="Portable output from generate_control_reproducibility_assets.py.",
    )
    parser.add_argument(
        "--drug-mosaic-plates",
        nargs="+",
        choices=drug.PLATES,
        default=list(drug.PLATES),
        help="Drug plates displayed in panel B; quantitative panels still use all plates.",
    )
    parser.add_argument("--output-dir", type=Path, default=OUTPUT_DIR)
    return parser.parse_args()


def contain_image(path: Path, size: tuple[int, int]) -> Image.Image:
    with Image.open(path) as source:
        return ImageOps.contain(source.convert("RGB"), size, Image.Resampling.LANCZOS)


def panel_font() -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(str(control.FONT_BOLD), control.FONT_PX)


def compose_figure(panel_paths: dict[str, Path], output_path: Path, *,
                   panel_boxes: dict | None = None,
                   canvas_size: tuple[int, int] = (CANVAS_WIDTH, CANVAS_HEIGHT),
                   footer_text: str | None = None) -> None:
    canvas = Image.new("RGB", canvas_size, "white")
    draw = ImageDraw.Draw(canvas)
    for label, box in (PANEL_BOXES if panel_boxes is None else panel_boxes).items():
        x, y, width, height = box
        panel = contain_image(panel_paths[label], (width, height))
        paste_x = x + (width - panel.width) // 2
        paste_y = y + (height - panel.height) // 2
        canvas.paste(panel, (paste_x, paste_y))
        draw.text((x - 28, y + 2), label, font=panel_font(), fill=INK, anchor="ra")
    if footer_text:
        font = ImageFont.truetype(str(control.FONT_REGULAR), 43)
        draw.text((canvas_size[0] / 2, canvas_size[1] - 50), footer_text,
                  font=font, fill=INK, anchor="mm")
    canvas.save(output_path, dpi=(FINAL_DPI, FINAL_DPI), compress_level=4)
    canvas.save(output_path.with_suffix(".pdf"), resolution=FINAL_DPI)


def add_picture_contained(slide, path: Path, box: tuple[int, int, int, int]) -> None:
    x_px, y_px, width_px, height_px = box
    scale = 1 / FINAL_DPI
    with Image.open(path) as image:
        image_aspect = image.width / image.height
    box_aspect = width_px / height_px
    if image_aspect > box_aspect:
        picture_width = width_px
        picture_height = width_px / image_aspect
        picture_x = x_px
        picture_y = y_px + (height_px - picture_height) / 2
    else:
        picture_height = height_px
        picture_width = height_px * image_aspect
        picture_x = x_px + (width_px - picture_width) / 2
        picture_y = y_px
    picture = slide.shapes.add_picture(
        str(path),
        Inches(picture_x * scale),
        Inches(picture_y * scale),
        width=Inches(picture_width * scale),
        height=Inches(picture_height * scale),
    )
    picture.name = f"Panel {path.stem.removeprefix('panel_')} - {path.name}"


def build_powerpoint(panel_paths: dict[str, Path], output_path: Path, *,
                     panel_boxes: dict | None = None,
                     canvas_size: tuple[int, int] = (CANVAS_WIDTH, CANVAS_HEIGHT),
                     footer_text: str | None = None) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(canvas_size[0] / FINAL_DPI)
    presentation.slide_height = Inches(canvas_size[1] / FINAL_DPI)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)

    for label, box in (PANEL_BOXES if panel_boxes is None else panel_boxes).items():
        add_picture_contained(slide, panel_paths[label], box)
        x, y, _, _ = box
        label_box = slide.shapes.add_textbox(
            Inches((x - 100) / FINAL_DPI),
            Inches((y + 2) / FINAL_DPI),
            Inches(0.20),
            Inches(0.22),
        )
        label_box.name = f"Panel label {label}"
        frame = label_box.text_frame
        frame.clear()
        frame.margin_left = 0
        frame.margin_right = 0
        frame.margin_top = 0
        frame.margin_bottom = 0
        paragraph = frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.RIGHT
        run = paragraph.add_run()
        run.text = label
        run.font.name = control.FONT_FAMILY
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = RGBColor(23, 33, 43)
    if footer_text:
        footer = slide.shapes.add_textbox(
            Inches(150 / FINAL_DPI), Inches((canvas_size[1] - 85) / FINAL_DPI),
            Inches((canvas_size[0] - 300) / FINAL_DPI), Inches(70 / FINAL_DPI),
        )
        footer.name = "Figure note"
        footer.text_frame.clear()
        footer.text_frame.margin_left = 0
        footer.text_frame.margin_right = 0
        footer.text_frame.margin_top = 0
        footer.text_frame.margin_bottom = 0
        paragraph = footer.text_frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.CENTER
        run = paragraph.add_run()
        run.text = footer_text
        run.font.name = control.FONT_FAMILY
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(23, 33, 43)
    presentation.save(output_path)


def build_panels(
    control_asset_root: Path,
    output_dir: Path,
    drug_mosaic_plates: tuple[str, ...],
) -> tuple[dict[str, Path], list[drug.Record], list[control.Record]]:
    panels_dir = output_dir / "panels"
    panels_dir.mkdir(parents=True, exist_ok=True)
    for stale_path in panels_dir.glob("panel_*.*"):
        if stale_path.stem.removeprefix("panel_") not in PANEL_BOXES:
            stale_path.unlink()
    paths = {label: panels_dir / f"panel_{label}.png" for label in PANEL_BOXES}

    drug_records = drug.read_records()
    drug.validate_inputs(drug_records)
    control_records = control.read_records(control_asset_root / "data.csv")
    control.validate_inputs(control_asset_root, control_records)

    drug.save_panel_a(paths["A"])
    drug.save_panel_b(
        drug_records,
        paths["B"],
        plates=drug_mosaic_plates,
    )
    for label, spec in zip("CDEF", drug.METRICS, strict=True):
        drug.save_metric_panel(drug_records, spec, paths[label], PANEL_BOXES[label][2:])
    drug_selected = drug.save_panel_g(drug_records, paths["G"], PANEL_BOXES["G"][2:])

    for label, spec in zip("HIJK", control.METRICS, strict=True):
        control.save_metric_panel(
            control_records,
            spec,
            paths[label],
            PANEL_BOXES[label][2:],
        )
    manifest = {
        "figure_scope": "drug-response assay plus no-drug inter-plate reproducibility",
        "ctb_measurement": "PRE CTB-647 channel 2 only",
        "drug_mosaic_plates": list(drug_mosaic_plates),
        "quantitative_drug_plates": list(drug.PLATES),
        "panel_roles": {
            "A": "shared plate and device context",
            "B": "drug-assay plate mosaics",
            "C-F": "drug-assay measurements",
            "G": "drug-assay representative overlays",
            "H-K": "no-drug reproducibility measurements",
        },
        "control_pre_exclusions": [record.key for record in control_records if record.pre_outlier],
        "drug_representatives": [record.key for record in drug_selected],
        "drug_data_sha256": drug.sha256(drug.DATA_CSV),
        "control_data_sha256": control.sha256(control_asset_root / "data.csv"),
        "font": "Arial 11 pt",
        "panels": {label: str(path) for label, path in paths.items()},
    }
    (output_dir / "figure_manifest.json").write_text(
        json.dumps(manifest, indent=2) + "\n",
        encoding="utf-8",
    )
    return paths, drug_records, control_records


def main() -> int:
    args = parse_args()
    output_dir = args.output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    drug_mosaic_plates = tuple(args.drug_mosaic_plates)
    panel_paths, _, _ = build_panels(
        args.control_asset_root.resolve(),
        output_dir,
        drug_mosaic_plates,
    )
    variant = f"{len(drug_mosaic_plates)}plate" + ("s" if len(drug_mosaic_plates) != 1 else "")
    figure_path = output_dir / f"axotomy_drug_and_control_supplementary_figure_{variant}.png"
    compose_figure(panel_paths, figure_path)
    build_powerpoint(
        panel_paths,
        output_dir / f"axotomy_drug_and_control_supplementary_figure_{variant}_editable.pptx",
    )
    print(f"Wrote {figure_path}")
    print(f"Wrote {figure_path.with_suffix('.pdf')}")
    print(
        "Wrote "
        + str(output_dir / f"axotomy_drug_and_control_supplementary_figure_{variant}_editable.pptx")
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
