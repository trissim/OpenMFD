#!/usr/bin/env python3

from __future__ import annotations

import csv
import hashlib
import json
import sys
from dataclasses import dataclass
from pathlib import Path
from statistics import mean, median, stdev

import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt
import numpy as np
from matplotlib import font_manager
from matplotlib.collections import LineCollection
from matplotlib.patches import FancyBboxPatch, Rectangle
from matplotlib.ticker import FuncFormatter
from PIL import Image, ImageDraw, ImageFont, ImageOps
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from axotomy_plate_mosaic import save_plate_mosaic

ROOT = Path(__file__).resolve().parents[3]
BIO_DIR = Path(__file__).resolve().parent
FIGURE_CODE_DIR = ROOT / "papers" / "mfd_platform" / "figures"
OUTPUT_DIR = BIO_DIR / "figure_output"
PANELS_DIR = OUTPUT_DIR / "panels"
DATA_CSV = BIO_DIR / "data.csv"
SOURCE_ARCHIVE = BIO_DIR / "OneDrive_1_8-24-2026.zip"

sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(FIGURE_CODE_DIR))

CONDITIONS = ("DMSO", "FC-A", "F04", "EpoB", "EpoB + FCA")
PLATES = ("p1", "p2", "p3", "p4")
FULL_PRE_AXON_INSET_PLATES = ("p1", "p2")
PLATE_COLORS = {
    "p1": "#0072B2",
    "p2": "#E69F00",
    "p3": "#009E73",
    "p4": "#CC79A7",
}
CONDITION_COLORS = {
    "DMSO": "#D8E2EA",
    "FC-A": "#C9E6DA",
    "F04": "#F8E2B3",
    "EpoB": "#EEC9CE",
    "EpoB + FCA": "#DDD3EA",
}
INK = "#17212B"
MUTED = "#5D6975"
GRID = "#D7DEE6"
EXCLUDED = "#D62828"
FONT_FAMILY = "Arial"
FINAL_DPI = 300
FONT_PX = round(11 * FINAL_DPI / 72)
FONT_REGULAR = Path(font_manager.findfont(font_manager.FontProperties(family=FONT_FAMILY)))
FONT_BOLD = Path(
    font_manager.findfont(font_manager.FontProperties(family=FONT_FAMILY, weight="bold"))
)

matplotlib.rcParams.update(
    {
        "font.family": FONT_FAMILY,
        "font.size": 11,
        "axes.titlesize": 11,
        "axes.labelsize": 11,
        "xtick.labelsize": 11,
        "ytick.labelsize": 11,
        "legend.fontsize": 11,
        "pdf.fonttype": 42,
        "ps.fonttype": 42,
    }
)

CANVAS_WIDTH = 5400
CANVAS_HEIGHT = 6520
PANEL_BOXES = {
    "A": (140, 100, 1390, 1370),
    "B": (1650, 100, 3610, 1370),
    "C": (140, 1600, 2510, 1660),
    "D": (2780, 1600, 2480, 1660),
    "E": (140, 3390, 2510, 1660),
    "F": (2780, 3390, 2480, 1660),
    "G": (140, 5200, 5120, 1160),
}


@dataclass(frozen=True)
class Record:
    plate: str
    well: str
    row: int
    column: int
    condition: str
    dapi_count: float
    pre_cell_count: float
    total_axon_count: float
    outlier: bool

    @property
    def key(self) -> str:
        return f"{self.plate}_{self.well}"

    @property
    def axon_per_pre(self) -> float:
        return self.total_axon_count / self.pre_cell_count


@dataclass(frozen=True)
class MetricSpec:
    panel: str
    title: str
    ylabel: str
    value: str


METRICS = (
    MetricSpec(
        "C",
        "Pre-axotomy CTB-positive cells",
        "CTB-647-positive somata (count)",
        "pre_cell_count",
    ),
    MetricSpec(
        "D",
        "Total axon measurement",
        "Total axon measurement (source units)",
        "total_axon_count",
    ),
    MetricSpec(
        "E",
        "Nuclear count",
        "Hoechst 33342-positive nuclei (count)",
        "dapi_count",
    ),
    MetricSpec(
        "F",
        "Axon measurement per pre-axotomy CTB-positive cell",
        "Total axon / pre-axotomy CTB-positive cell",
        "axon_per_pre",
    ),
)


def read_records() -> list[Record]:
    records = []
    with DATA_CSV.open(newline="", encoding="utf-8") as handle:
        for row in csv.DictReader(handle):
            records.append(
                Record(
                    plate=row["plate"],
                    well=row["well"],
                    row=int(row["row"].removeprefix("R")),
                    column=int(row["col"]),
                    condition=row["condition"],
                    dapi_count=float(row["dapi_count"]),
                    pre_cell_count=float(row["pre_cell_count"]),
                    total_axon_count=float(row["total_axon_count"]),
                    outlier=row["outlier"].strip().lower() == "true",
                )
            )
    return records


def image_path(record: Record, overlay: bool = True) -> Path:
    variant = "overlay" if overlay else "no_overlay"
    suffix = "overlay" if overlay else "no_overlay"
    matches = tuple(
        (BIO_DIR / record.plate / variant).glob(f"{record.plate}_*_{record.well}*_{suffix}.tiff")
    )
    if len(matches) != 1:
        raise ValueError(f"Expected one {variant} image for {record.key}; found {matches}")
    return matches[0]


def validate_inputs(records: list[Record]) -> None:
    if len(records) != 96:
        raise ValueError(f"Expected 96 data rows; found {len(records)}")
    if len({record.key for record in records}) != len(records):
        raise ValueError("Duplicate plate/well keys in data.csv")
    if {record.plate for record in records} != set(PLATES):
        raise ValueError("Unexpected plate identifiers in data.csv")
    if {record.condition for record in records} != set(CONDITIONS):
        raise ValueError("Unexpected condition names in data.csv")

    for plate in PLATES:
        plate_records = [record for record in records if record.plate == plate]
        if len(plate_records) != 24:
            raise ValueError(f"Expected 24 devices for {plate}; found {len(plate_records)}")
        coordinates = {(record.row, record.column) for record in plate_records}
        expected = {(row, column) for row in range(2, 8) for column in range(2, 6)}
        if coordinates != expected:
            raise ValueError(f"Unexpected acquired coordinates for {plate}")

    for record in records:
        for overlay in (True, False):
            path = image_path(record, overlay=overlay)
            marker = "_outlier_" in path.name
            if marker != record.outlier:
                raise ValueError(f"Outlier marker mismatch for {record.key}: {path.name}")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while chunk := handle.read(1024 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


def pil_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(str(FONT_BOLD if bold else FONT_REGULAR), size)


def save_panel_a(path: Path) -> None:
    # Graph-only figures do not need the CAD dependencies or design assets.
    from generate_openmfd_design_figure import (
        COLORS as OPENMFD_COLORS, FRAME_STL, Bounds, DxfSources, StlMesh,
        add_bounds_background, draw_dxf_layer_raw, setup_raw_dxf_axis,
    )
    from openmfd.devices.presets import TwoCompartmentDeviceConfig

    config = TwoCompartmentDeviceConfig()
    dxf = DxfSources.from_generated_design()
    frame = StlMesh.from_file(FRAME_STL)
    frame_vertices = frame.triangles.reshape((-1, 3))
    plate_width = float(np.ptp(frame_vertices[:, 0]))
    plate_height = float(np.ptp(frame_vertices[:, 1]))
    array_width = config.grid_size[0] * config.casing_x
    array_height = config.grid_size[1] * config.casing_y
    array_left = (plate_width - array_width) / 2
    array_bottom = (plate_height - array_height) / 2

    fig = plt.figure(figsize=(7.0, 7.0), facecolor="white")
    fig.suptitle(
        "OpenMFD plate and assay device",
        fontsize=18,
        fontweight="bold",
        color=INK,
        y=0.98,
    )
    plate_ax = fig.add_axes((0.05, 0.45, 0.90, 0.45))
    plate_ax.set_xlim(-2, plate_width + 2)
    plate_ax.set_ylim(-2, plate_height + 2)
    plate_ax.set_aspect("equal")
    plate_ax.axis("off")
    plate_ax.add_patch(
        FancyBboxPatch(
            (0, 0),
            plate_width,
            plate_height,
            boxstyle="round,pad=0,rounding_size=4",
            facecolor="#F4F7FA",
            edgecolor=INK,
            linewidth=1.4,
        )
    )

    interior_left = array_left + config.casing_x
    interior_bottom = array_bottom + config.casing_y
    plate_ax.add_patch(
        Rectangle(
            (interior_left, interior_bottom),
            4 * config.casing_x,
            6 * config.casing_y,
            facecolor="#FFF1B8",
            edgecolor="#D39A00",
            linewidth=1.4,
            alpha=0.65,
            zorder=1,
        )
    )

    top_segments = []
    bottom_segments = []
    for x_index in range(config.grid_size[0]):
        for y_index in range(config.grid_size[1]):
            dx = array_left + x_index * config.casing_x
            dy = array_bottom + y_index * config.casing_y
            top_segments.extend(
                [
                    ((a[0] + dx, a[1] + dy), (b[0] + dx, b[1] + dy))
                    for a, b in dxf.single_top.segments
                ]
            )
            bottom_segments.extend(
                [
                    ((a[0] + dx, a[1] + dy), (b[0] + dx, b[1] + dy))
                    for a, b in dxf.single_bottom.segments
                ]
            )
    plate_ax.add_collection(
        LineCollection(
            bottom_segments, colors=OPENMFD_COLORS["bottom"], linewidths=0.28, alpha=0.70, zorder=2
        )
    )
    plate_ax.add_collection(
        LineCollection(
            top_segments, colors=OPENMFD_COLORS["top"], linewidths=0.52, alpha=0.96, zorder=3
        )
    )
    plate_ax.text(
        plate_width / 2,
        -1.0,
        "48 devices / 96 wells",
        ha="center",
        va="top",
        color=INK,
        fontsize=12,
        fontweight="bold",
    )
    plate_ax.text(
        interior_left + 2 * config.casing_x,
        interior_bottom + 3 * config.casing_y,
        "24-device interior\nimaging region",
        ha="center",
        va="center",
        color="#815E00",
        fontsize=10,
        fontweight="bold",
        bbox={"facecolor": "white", "edgecolor": "none", "alpha": 0.82, "pad": 2.5},
        zorder=4,
    )

    zoom_ax = fig.add_axes((0.12, 0.07, 0.76, 0.28))
    zoom_bounds = Bounds.combine((dxf.single_top, dxf.single_bottom))
    setup_raw_dxf_axis(zoom_ax, zoom_bounds, 0.7, 0.7)
    add_bounds_background(zoom_ax, zoom_bounds, 0.7, 0.7)
    draw_dxf_layer_raw(zoom_ax, dxf.single_bottom, OPENMFD_COLORS["bottom"], 0.55, 0.72)
    draw_dxf_layer_raw(zoom_ax, dxf.single_top, OPENMFD_COLORS["top"], 1.5, 1.0)
    zoom_ax.text(4.5, 1.63, "Soma", ha="center", va="top", fontsize=11, color=INK)
    zoom_ax.text(13.5, 1.63, "Distal", ha="center", va="top", fontsize=11, color=INK)
    zoom_ax.text(
        9.0,
        8.05,
        "Two-compartment device",
        ha="center",
        va="bottom",
        fontsize=13,
        fontweight="bold",
        color=INK,
    )
    zoom_ax.text(
        9.0, 0.95, "300 um microchannel barrier", ha="center", va="top", fontsize=10, color=MUTED
    )

    fig.savefig(path, dpi=300, facecolor="white")
    fig.savefig(path.with_suffix(".svg"), facecolor="white")
    plt.close(fig)


def image_thumbnail(
    path: Path,
    size: tuple[int, int],
    background: str = "white",
    crop_top_fraction: float = 0.0,
) -> Image.Image:
    with Image.open(path) as source:
        image = source.convert("RGB")
        if crop_top_fraction:
            image = image.crop(
                (0, round(image.height * crop_top_fraction), image.width, image.height)
            )
        image.thumbnail(size, Image.Resampling.LANCZOS, reducing_gap=3.0)
    canvas = Image.new("RGB", size, background)
    canvas.paste(image, ((size[0] - image.width) // 2, (size[1] - image.height) // 2))
    return canvas


def save_panel_b(
    records: list[Record],
    path: Path,
    plates: tuple[str, ...] = PLATES,
) -> None:
    if not plates or not set(plates).issubset(PLATES):
        raise ValueError(f"Panel B plates must be selected from {PLATES}: {plates}")
    if plates == PLATES:
        title = "All drug-plate device overlays in plate coordinates"
    elif len(plates) == 1:
        title = f"Device overlays for Drug Plate {plates[0][1:]} in plate coordinates"
    else:
        plate_names = ", ".join(plate[1:] for plate in plates[:-1])
        title = (
            f"Device overlays for Drug Plates {plate_names} and {plates[-1][1:]} "
            "in plate coordinates"
        )
    save_plate_mosaic(
        records,
        path,
        plates=plates,
        image_path_for_record=image_path,
        is_excluded=lambda record: record.outlier,
        plate_colors=PLATE_COLORS,
        legend_plates=PLATES,
        title=title,
        font_regular=FONT_REGULAR,
        font_bold=FONT_BOLD,
    )


def metric_value(record: Record, name: str) -> float:
    return float(getattr(record, name))


def point_x(record: Record, condition_index: int) -> float:
    plate_offset = {"p1": -0.24, "p2": -0.08, "p3": 0.08, "p4": 0.24}[record.plate]
    deterministic_jitter = (((record.row * 7 + record.column * 11) % 7) - 3) * 0.012
    return condition_index + plate_offset + deterministic_jitter


def save_metric_panel(
    records: list[Record],
    spec: MetricSpec,
    path: Path,
    size: tuple[int, int] | None = None,
    *,
    conditions: tuple[str, ...] = CONDITIONS,
    plates: tuple[str, ...] = PLATES,
    condition_colors: dict[str, str] = CONDITION_COLORS,
    plate_colors: dict[str, str] = PLATE_COLORS,
    condition_labels: dict[str, str] | None = None,
    footer: str = "Bars: mean +/- SD; points: retained devices",
) -> None:
    retained = [record for record in records if not record.outlier]
    width, height = size or PANEL_BOXES[spec.panel][2:]
    figure, axis = plt.subplots(
        figsize=(width / FINAL_DPI, height / FINAL_DPI),
        dpi=FINAL_DPI,
        facecolor="white",
    )
    x_values = np.arange(len(conditions), dtype=float)
    means = []
    deviations = []
    counts = []
    for condition in conditions:
        values = [
            metric_value(record, spec.value) for record in retained if record.condition == condition
        ]
        means.append(mean(values) if values else np.nan)
        deviations.append(stdev(values) if len(values) > 1 else 0.0)
        counts.append(len(values))

    axis.bar(
        x_values,
        means,
        width=0.68,
        color=[condition_colors[condition] for condition in conditions],
        edgecolor=INK,
        linewidth=1.1,
        zorder=1,
    )
    lower = np.minimum(np.asarray(means), np.asarray(deviations))
    axis.errorbar(
        x_values,
        means,
        yerr=np.vstack((lower, deviations)),
        fmt="none",
        ecolor=INK,
        elinewidth=1.2,
        capsize=4,
        capthick=1.2,
        zorder=2,
    )
    offsets = dict(zip(plates, np.linspace(-0.24, 0.24, len(plates))))
    for plate in plates:
        plate_records = [record for record in retained if record.plate == plate]
        axis.scatter(
            [conditions.index(record.condition) + offsets[plate]
             + (((record.row * 7 + record.column * 11) % 7) - 3) * 0.012
             for record in plate_records],
            [metric_value(record, spec.value) for record in plate_records],
            s=31,
            color=plate_colors[plate],
            edgecolors="white",
            linewidths=0.55,
            alpha=0.92,
            label=f"Plate {plate[1:]}",
            zorder=3,
        )

    axis.set_title(spec.title, fontsize=11, color=INK, pad=32)
    axis.set_ylabel(spec.ylabel, fontsize=11, color=INK)
    axis.set_xticks(x_values)
    axis.set_xticklabels(
        [f"{(condition_labels or {}).get(condition, condition)}\nn={count}"
         for condition, count in zip(conditions, counts)],
        fontsize=11,
    )
    axis.set_xlim(-0.55, len(conditions) - 0.45)
    axis.set_ylim(bottom=0)
    axis.yaxis.set_major_formatter(FuncFormatter(lambda value, _position: f"{value:,.0f}"))
    axis.grid(axis="y", color=GRID, linewidth=0.8, alpha=0.85, zorder=0)
    axis.spines[["top", "right"]].set_visible(False)
    axis.spines[["left", "bottom"]].set_color("#7A8793")
    axis.tick_params(colors=INK)
    axis.legend(
        loc="lower center",
        bbox_to_anchor=(0.5, 1.01),
        ncol=4,
        frameon=False,
        fontsize=11,
        handletextpad=0.35,
        columnspacing=1.2,
    )
    figure.text(
        0.99,
        0.01,
        footer,
        ha="right",
        va="bottom",
        fontsize=11,
        color=MUTED,
    )
    figure.tight_layout(rect=(0.03, 0.05, 0.995, 0.96))
    figure.savefig(path, dpi=FINAL_DPI, facecolor="white")
    figure.savefig(path.with_suffix(".svg"), facecolor="white")
    plt.close(figure)


def representative_records(
    records: list[Record],
    eligible_plates: tuple[str, ...] | None = None,
) -> list[Record]:
    selected = []
    for condition in CONDITIONS:
        candidates = [
            record
            for record in records
            if record.condition == condition
            and not record.outlier
            and (eligible_plates is None or record.plate in eligible_plates)
        ]
        if not candidates:
            raise ValueError(
                f"No eligible representative devices for {condition} in {eligible_plates}"
            )
        vectors = {
            record.key: {
                "pre_cell_count": record.pre_cell_count,
                "total_axon_count": record.total_axon_count,
                "dapi_count": record.dapi_count,
                "axon_per_pre": record.axon_per_pre,
            }
            for record in candidates
        }
        metric_names = tuple(next(iter(vectors.values())).keys())
        medians = {
            name: median(vector[name] for vector in vectors.values()) for name in metric_names
        }
        mads = {
            name: median(abs(vector[name] - medians[name]) for vector in vectors.values()) or 1.0
            for name in metric_names
        }

        scores = {
            record.key: sum(
                ((vectors[record.key][name] - medians[name]) / mads[name]) ** 2
                for name in metric_names
            )
            for record in candidates
        }
        selected.append(min(candidates, key=lambda record: scores[record.key]))
    return selected


def save_panel_g(
    records: list[Record],
    path: Path,
    size: tuple[int, int] | None = None,
) -> list[Record]:
    selected = representative_records(records, eligible_plates=FULL_PRE_AXON_INSET_PLATES)
    canvas_width, canvas_height = size or PANEL_BOXES["G"][2:]
    gap = 32
    top = 125
    bottom = 60
    tile_width = (canvas_width - (len(selected) - 1) * gap) // len(selected)
    tile_height = canvas_height - top - bottom
    canvas = Image.new("RGB", (canvas_width, canvas_height), "white")
    draw = ImageDraw.Draw(canvas)
    draw.text(
        (canvas_width / 2, 5),
        "Representative non-outlier device overlays",
        font=pil_font(FONT_PX, bold=True),
        fill=INK,
        anchor="ma",
    )
    for index, record in enumerate(selected):
        x = index * (tile_width + gap)
        tile = image_thumbnail(
            image_path(record),
            (tile_width, tile_height),
            crop_top_fraction=0.05,
        )
        canvas.paste(tile, (x, top))
        draw.text(
            (x + tile_width / 2, 64),
            record.condition,
            font=pil_font(FONT_PX, bold=True),
            fill=INK,
            anchor="ma",
        )
        draw.text(
            (x + tile_width / 2, canvas_height - 14),
            record.plate.upper(),
            font=pil_font(FONT_PX),
            fill=MUTED,
            anchor="ms",
        )
    canvas.save(path, dpi=(300, 300), compress_level=5)
    return selected


def contain_image(path: Path, size: tuple[int, int]) -> Image.Image:
    with Image.open(path) as source:
        return ImageOps.contain(source.convert("RGB"), size, Image.Resampling.LANCZOS)


def compose_figure(panel_paths: dict[str, Path], output_path: Path) -> None:
    canvas = Image.new("RGB", (CANVAS_WIDTH, CANVAS_HEIGHT), "white")
    draw = ImageDraw.Draw(canvas)
    label_font = pil_font(86, bold=True)
    for label, box in PANEL_BOXES.items():
        x, y, width, height = box
        panel = contain_image(panel_paths[label], (width, height))
        paste_x = x + (width - panel.width) // 2
        paste_y = y + (height - panel.height) // 2
        canvas.paste(panel, (paste_x, paste_y))
        draw.text((x - 22, y + 4), label, font=label_font, fill=INK, anchor="ra")
    canvas.save(output_path, dpi=(300, 300), compress_level=4)
    canvas.save(output_path.with_suffix(".pdf"), resolution=300.0)


def add_picture_contained(
    slide, path: Path, box: tuple[float, float, float, float], slide_width: float
):
    x_px, y_px, width_px, height_px = box
    x = x_px / CANVAS_WIDTH * slide_width
    y = y_px / CANVAS_WIDTH * slide_width
    width = width_px / CANVAS_WIDTH * slide_width
    height = height_px / CANVAS_WIDTH * slide_width
    with Image.open(path) as image:
        image_aspect = image.width / image.height
    box_aspect = width / height
    if image_aspect > box_aspect:
        picture_width = width
        picture_height = width / image_aspect
        picture_x = x
        picture_y = y + (height - picture_height) / 2
    else:
        picture_height = height
        picture_width = height * image_aspect
        picture_x = x + (width - picture_width) / 2
        picture_y = y
    return slide.shapes.add_picture(
        str(path),
        Inches(picture_x),
        Inches(picture_y),
        width=Inches(picture_width),
        height=Inches(picture_height),
    )


def build_powerpoint(panel_paths: dict[str, Path], output_path: Path) -> None:
    presentation = Presentation()
    slide_width = 13.333
    slide_height = slide_width * CANVAS_HEIGHT / CANVAS_WIDTH
    presentation.slide_width = Inches(slide_width)
    presentation.slide_height = Inches(slide_height)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = RGBColor(255, 255, 255)

    for label, box in PANEL_BOXES.items():
        picture = add_picture_contained(slide, panel_paths[label], box, slide_width)
        picture.name = f"Panel {label} - {panel_paths[label].name}"
        x, y, _width, _height = box
        label_box = slide.shapes.add_textbox(
            Inches((x - 112) / CANVAS_WIDTH * slide_width),
            Inches((y + 3) / CANVAS_WIDTH * slide_width),
            Inches(0.25),
            Inches(0.28),
        )
        label_box.name = f"Panel label {label}"
        frame = label_box.text_frame
        frame.clear()
        frame.margin_left = 0
        frame.margin_right = 0
        frame.margin_top = 0
        frame.margin_bottom = 0
        paragraph = frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.RIGHT
        run = paragraph.add_run()
        run.text = label
        run.font.name = "Arial"
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = RGBColor(23, 33, 43)

    presentation.save(output_path)


def write_manifest(
    records: list[Record], selected: list[Record], panel_paths: dict[str, Path]
) -> None:
    retained = [record for record in records if not record.outlier]
    manifest = {
        "authoritative_data": str(DATA_CSV.relative_to(ROOT)),
        "data_sha256": sha256(DATA_CSV),
        "source_archive": str(SOURCE_ARCHIVE.relative_to(ROOT)),
        "records_total": len(records),
        "records_retained": len(retained),
        "records_excluded": len(records) - len(retained),
        "retained_by_condition": {
            condition: sum(record.condition == condition for record in retained)
            for condition in CONDITIONS
        },
        "panels": {label: str(path.relative_to(ROOT)) for label, path in panel_paths.items()},
        "representative_devices": {record.condition: record.key for record in selected},
        "representative_rule": (
            "Minimum squared median-absolute-deviation distance across pre_cell_count, "
            "total_axon_count, dapi_count, and total_axon_count/pre_cell_count among "
            "non-outlier devices from Plates 1 and 2. These plates contain the full "
            "PRE stitched CTB image required for a consistent axon-compartment inset."
        ),
    }
    (OUTPUT_DIR / "figure_manifest.json").write_text(
        json.dumps(manifest, indent=2) + "\n", encoding="utf-8"
    )


def main() -> int:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    PANELS_DIR.mkdir(parents=True, exist_ok=True)
    records = read_records()
    validate_inputs(records)

    panel_paths = {label: PANELS_DIR / f"panel_{label}.png" for label in "ABCDEFG"}
    save_panel_a(panel_paths["A"])
    save_panel_b(records, panel_paths["B"])
    for spec in METRICS:
        save_metric_panel(records, spec, panel_paths[spec.panel])
    selected = save_panel_g(records, panel_paths["G"])

    figure_path = OUTPUT_DIR / "axotomy_supplementary_figure.png"
    compose_figure(panel_paths, figure_path)
    build_powerpoint(
        panel_paths,
        OUTPUT_DIR / "axotomy_supplementary_figure_editable.pptx",
    )
    write_manifest(records, selected, panel_paths)

    print(f"Wrote {figure_path}")
    print(f"Wrote {figure_path.with_suffix('.pdf')}")
    print(f"Wrote {OUTPUT_DIR / 'axotomy_supplementary_figure_editable.pptx'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
